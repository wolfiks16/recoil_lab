"""Excel-экспорт результатов теплового сценария (ThermalRun)."""

from __future__ import annotations

from io import BytesIO

import openpyxl
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def _header_row(ws, headers: list[str]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)


def _autosize(ws, max_width: int = 35) -> None:
    for col_cells in ws.columns:
        max_len = 0
        col_letter = col_cells[0].column_letter
        for cell in col_cells:
            try:
                max_len = max(max_len, len("" if cell.value is None else str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 2, max_width)


def export_thermal_results_to_excel(thermal_run) -> BytesIO:
    """Генерирует XLSX-отчёт по тепловому сценарию и возвращает BytesIO."""
    config = thermal_run.config_snapshot or {}
    result = thermal_run.result_snapshot or {}

    timeline = result.get("timeline") or {}
    peaks = result.get("peaks") or {}
    cycles = result.get("cycles") or []
    node_names = result.get("node_names") or []
    node_display_names = result.get("node_display_names") or []
    heats_total_j = result.get("heats_total_j") or []

    network = config.get("network") or {}
    nodes_table = network.get("nodes") or []
    links_table = network.get("links") or []
    preset = config.get("preset") or thermal_run.network_preset

    wb = openpyxl.Workbook()

    # ── Лист 1: Summary ──────────────────────────────────────────────────────
    ws_sum = wb.active
    ws_sum.title = "summary"
    _header_row(ws_sum, ["Параметр", "Значение"])

    ws_sum.append(["Имя сценария", thermal_run.name])
    ws_sum.append(["Пресет сети", preset])
    ws_sum.append(["Число циклов", thermal_run.repetitions])
    ws_sum.append(["Пауза, с", float(thermal_run.pause_s)])
    ws_sum.append(["T_max, °C", float(thermal_run.max_temp_c) if thermal_run.max_temp_c is not None else None])
    ws_sum.append(["Узел T_max", thermal_run.max_temp_node_name or ""])
    ws_sum.append(["Общее тепло, Дж", float(thermal_run.total_heat_j) if thermal_run.total_heat_j is not None else None])
    for i, q in enumerate(heats_total_j):
        ws_sum.append([f"Тепло тормоза {i + 1}, Дж", float(q)])
    if thermal_run.warnings_text:
        ws_sum.append(["Предупреждения", thermal_run.warnings_text])
    ws_sum.append(["Создан", str(thermal_run.created_at)])

    _autosize(ws_sum)

    # ── Лист 2: Пики по узлам ────────────────────────────────────────────────
    if peaks:
        ws_pk = wb.create_sheet("peaks")
        _header_row(ws_pk, ["Узел", "T_пик, °C", "t, с", "Цикл", "Сегмент"])
        for key, p in peaks.items():
            ws_pk.append([
                p.get("display_name") or key,
                p.get("temp_peak_c"),
                p.get("t_peak_s"),
                p.get("cycle_at_peak"),
                p.get("segment_at_peak") or "",
            ])
        _autosize(ws_pk)

    # ── Лист 3: Timeline (t, температуры, мощности) + scatter-график ─────────
    t_list = timeline.get("t") or []
    temp_nodes = timeline.get("temp_nodes") or []    # [[T_node0, T_node1, …], …]
    power_brakes = timeline.get("power_brakes") or []  # [[P_brake0, …], …]
    heat_brakes = timeline.get("heat_brakes") or []   # [[Q_brake0, …], …]
    cycle_idx = timeline.get("cycle_index") or []
    segment = timeline.get("segment") or []

    if t_list:
        n_nodes = len(node_names)
        n_brakes = len(power_brakes[0]) if power_brakes else 0
        has_heat = bool(heat_brakes)

        display = [node_display_names[i] if i < len(node_display_names) else node_names[i] for i in range(n_nodes)]
        tl_headers = (
            ["t, с"]
            + [f"T {d}, °C" for d in display]
            + [f"P тормоза {j + 1}, Вт" for j in range(n_brakes)]
            + ([f"Q тормоза {j + 1}, кДж" for j in range(n_brakes)] if has_heat else [])
            + ["Цикл", "Сегмент"]
        )

        ws_tl = wb.create_sheet("timeline")
        _header_row(ws_tl, tl_headers)

        for i, t in enumerate(t_list):
            row: list = [t]
            row.extend(temp_nodes[i] if i < len(temp_nodes) else [None] * n_nodes)
            row.extend(power_brakes[i] if i < len(power_brakes) else [None] * n_brakes)
            if has_heat:
                raw = heat_brakes[i] if i < len(heat_brakes) else [None] * n_brakes
                row.extend([v / 1000.0 if v is not None else None for v in raw])
            row.append(cycle_idx[i] if i < len(cycle_idx) else None)
            row.append(segment[i] if i < len(segment) else "")
            ws_tl.append(row)

        _autosize(ws_tl, max_width=20)

        if n_nodes > 0:
            chart = ScatterChart()
            chart.title = "T(t) — температуры узлов"
            chart.x_axis.title = "t, с"
            chart.y_axis.title = "T, °C"
            chart.style = 2
            chart.height = 16
            chart.width = 28
            xvals = Reference(ws_tl, min_col=1, min_row=2, max_row=ws_tl.max_row)
            for col in range(2, 2 + n_nodes):
                yvals = Reference(ws_tl, min_col=col, min_row=1, max_row=ws_tl.max_row)
                chart.series.append(Series(yvals, xvals, title_from_data=True))
            anchor = get_column_letter(len(tl_headers) + 2) + "2"
            ws_tl.add_chart(chart, anchor)

    # ── Лист 4: Сводка по циклам ─────────────────────────────────────────────
    if cycles:
        ws_cyc = wb.create_sheet("cycles")
        cyc_headers = ["Цикл"] + [f"T_max {node_names[i]}, °C" for i in range(len(node_names))] + ["Тепло по тормозам, Дж"]
        _header_row(ws_cyc, cyc_headers)
        for c in cycles:
            max_temps = c.get("max_temp_in_cycle_c") or {}
            added_heat = c.get("added_heat_by_brake_j") or {}
            row = [c.get("cycle_number")]
            for name in node_names:
                row.append(max_temps.get(name))
            row.append(" · ".join(f"{float(v):.2f}" for v in added_heat.values()))
            ws_cyc.append(row)
        _autosize(ws_cyc)

    # ── Лист 5: Узлы сети ────────────────────────────────────────────────────
    if nodes_table:
        ws_nd = wb.create_sheet("nodes")
        if preset == "user_simple":
            _header_row(ws_nd, ["Узел", "m, кг", "cp, Дж/(кг·К)", "C, Дж/К", "G_воздух, Вт/К", "T₀, °C"])
            for nd in nodes_table:
                h = float(nd.get("h_ambient_w_per_m2k") or 0.0)
                a = float(nd.get("area_ambient_m2") or 0.0)
                ws_nd.append([
                    nd.get("display_name") or nd.get("name"),
                    nd.get("mass_kg"),
                    nd.get("cp_j_per_kgk"),
                    nd.get("capacitance_j_per_k"),
                    round(h * a, 6),
                    nd.get("temp0_c"),
                ])
        else:
            _header_row(ws_nd, ["Узел", "Материал", "m, кг", "cp, Дж/(кг·К)", "C, Дж/К",
                                 "A_amb, м²", "h_amb, Вт/(м²·К)", "A_rad, м²", "T₀, °C"])
            for nd in nodes_table:
                ws_nd.append([
                    nd.get("display_name") or nd.get("name"),
                    nd.get("material_key") or "",
                    nd.get("mass_kg"),
                    nd.get("cp_j_per_kgk"),
                    nd.get("capacitance_j_per_k"),
                    nd.get("area_ambient_m2"),
                    nd.get("h_ambient_w_per_m2k"),
                    nd.get("area_radiation_m2"),
                    nd.get("temp0_c"),
                ])
        _autosize(ws_nd)

    # ── Лист 6: Связи сети ───────────────────────────────────────────────────
    if links_table:
        ws_lk = wb.create_sheet("links")
        if preset == "user_simple":
            _header_row(ws_lk, ["A", "B", "G, Вт/К", "Описание"])
            for lk in links_table:
                ws_lk.append([
                    lk.get("node_a"),
                    lk.get("node_b"),
                    lk.get("conductance_w_per_k"),
                    lk.get("description") or "",
                ])
        else:
            _header_row(ws_lk, ["A", "B", "h, Вт/(м²·К)", "A, м²", "G, Вт/К", "Описание"])
            for lk in links_table:
                ws_lk.append([
                    lk.get("node_a"),
                    lk.get("node_b"),
                    lk.get("h_w_per_m2k"),
                    lk.get("area_m2"),
                    lk.get("conductance_w_per_k"),
                    lk.get("description") or "",
                ])
        _autosize(ws_lk)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
