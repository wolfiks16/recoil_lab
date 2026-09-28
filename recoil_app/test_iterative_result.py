"""Срез 4: этапы итерационного расчёта на итоговом расчёте — графики, XLSX, таблицы, страница."""

from __future__ import annotations

import shutil
import tempfile
from dataclasses import replace
from pathlib import Path

import numpy as np
import openpyxl
from django.contrib.auth import get_user_model
from django.test import SimpleTestCase, TestCase, override_settings
from django.urls import reverse

from .services.charting import save_interactive_charts
from .services.dynamics import simulate_recoil_core
from .services.iterative import MODE_RECOIL, IterativeSession
from .services.iterative import store
from .services.iterative.overlay import build_stage_overlay
from .services.iterative.result_view import build_stage_tables
from .services.reporting import export_results_to_excel
from .test_iterative import B1, B2, CURVE, DRIVE, RECOIL
from .test_iterative_store import _drive_xlsx


def _staged_outcome():
    """Откат с двумя сменами: bz тормоза 1 на 0.25 м, затем выключение 2-го и новый табличный на 0.5 м."""
    session = IterativeSession(MODE_RECOIL, RECOIL, [B1, B2], DRIVE)
    session.advance_distance(0.25)
    session.reconfigure([replace(B1, bz=0.6), B2])
    session.advance_distance(0.25)
    session.reconfigure([replace(B1, bz=0.6), None, CURVE])
    session.run_to_end()
    return session.build_result()


class OverlayTests(SimpleTestCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.outcome = _staged_outcome()
        cls.overlay = build_stage_overlay(cls.outcome)

    def test_segments_cover_timeline_in_stage_order(self):
        segments = self.overlay["segments"]
        self.assertEqual([s["stage"] for s in segments], [0, 1, 2, 1, 0])
        t = self.outcome.result.t
        self.assertEqual(segments[0]["t0"], float(t[0]))
        self.assertEqual(segments[-1]["t1"], float(t[-1]))
        for a, b in zip(segments, segments[1:]):
            self.assertEqual(a["t1"], b["t0"])            # узел-дубль: стык ровно в момент переключения

    def test_x_bands_follow_switch_points(self):
        bands = self.overlay["x_bands"]
        self.assertEqual([b["stage"] for b in bands], [1, 2])
        self.assertAlmostEqual(bands[0]["x0"], 0.25, delta=1e-9)
        self.assertAlmostEqual(bands[0]["x1"], 0.50, delta=1e-9)
        self.assertAlmostEqual(bands[1]["x1"], float(np.max(self.outcome.result.x)))

    def test_stage_rows_describe_changes(self):
        stages = self.overlay["stages"]
        self.assertEqual(stages[1]["changes"], ["тормоз 1: bz 0.83→0.6"])
        self.assertEqual(stages[2]["brakes"][1]["kind"], "off")
        self.assertEqual(stages[2]["brakes"][2]["kind"], "curve")
        self.assertEqual(len(stages[2]["brakes"][2]["points"]), 5)

    def test_charts_and_xlsx_include_stages(self):
        with tempfile.TemporaryDirectory() as tmp:
            files = save_interactive_charts(self.outcome.result, tmp, prefix="it", stage_overlay=self.overlay)
            annotated = Path(files["chart_x_t_annotated"]).read_text(encoding="utf-8")
            self.assertIn("этап 1", annotated)
            # Этапы помечены и скрыты по умолчанию — показывает переключатель на странице.
            self.assertIn('"name":"stage-band"', annotated)
            self.assertIn('"name":"stage-label"', annotated)
            self.assertNotIn('"visible":true', annotated)
            self.assertIn("этап 2", Path(files["chart_v_x"]).read_text(encoding="utf-8"))
            self.assertIn("Σ тормозов · этап 2", Path(files["chart_fmag_v"]).read_text(encoding="utf-8"))
            self.assertIn("этап 1", Path(files["chart_x_t_return"]).read_text(encoding="utf-8"))

            path = Path(tmp) / "it.xlsx"
            export_results_to_excel(self.outcome.result, path, stage_overlay=self.overlay)
            wb = openpyxl.load_workbook(path, read_only=True)
            self.assertIn("Этапы", wb.sheetnames)
            self.assertIn("Этапы_тормоза", wb.sheetnames)
            data = wb["data"]
            header = next(data.iter_rows(min_row=1, max_row=1, values_only=True))
            self.assertEqual(header[-1], "этап")
            stages_sheet = list(wb["Этапы"].iter_rows(values_only=True))
            self.assertEqual(len(stages_sheet), 1 + 3)
            self.assertEqual(len(list(wb["Этапы_тормоза"].iter_rows(values_only=True))), 1 + 3 * 3)
            wb.close()

    def test_plain_run_has_no_stage_artifacts(self):
        result = simulate_recoil_core(*DRIVE, RECOIL, [B1, B2])
        with tempfile.TemporaryDirectory() as tmp:
            files = save_interactive_charts(result, tmp, prefix="plain")
            self.assertNotIn("этап", Path(files["chart_x_t_annotated"]).read_text(encoding="utf-8"))
            path = Path(tmp) / "plain.xlsx"
            export_results_to_excel(result, path)
            wb = openpyxl.load_workbook(path, read_only=True)
            self.assertNotIn("Этапы", wb.sheetnames)
            header = next(wb["data"].iter_rows(min_row=1, max_row=1, values_only=True))
            self.assertNotEqual(header[-1], "этап")
            wb.close()


class ResultPageTests(TestCase):

    @classmethod
    def setUpClass(cls):
        cls._media = tempfile.mkdtemp(prefix="iter_result_media_")
        cls._override = override_settings(MEDIA_ROOT=cls._media)
        cls._override.enable()
        super().setUpClass()

    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        cls._override.disable()
        shutil.rmtree(cls._media, ignore_errors=True)

    def setUp(self):
        for child in Path(self._media).iterdir():
            shutil.rmtree(child, ignore_errors=True)
        self.user = get_user_model().objects.create_superuser("res_admin", "r@example.com", "pw")
        self.client.force_login(self.user)
        calc = store.create_calc(name="it_result", mode=MODE_RECOIL, recoil=RECOIL, config=[B1, B2],
                                 owner=self.user, input_file=_drive_xlsx())
        store.advance_distance(calc, 0.25)
        store.reconfigure(calc, [replace(B1, bz=0.6), B2])
        self.calc = calc
        self.run = store.finish(calc)

    def test_stage_tables(self):
        tables = build_stage_tables(self.run)
        self.assertEqual([s["index"] for s in tables["stages"]], [0, 1])
        brake1 = tables["brakes"][0]
        self.assertEqual(brake1["changed_params"], ["B̄₃"])
        bz_row = next(r for r in brake1["rows"] if r["label"] == "B̄₃")
        self.assertEqual([c["changed"] for c in bz_row["cells"]], [False, True])
        self.assertEqual(tables["brakes"][1]["changed_params"], [])

    def test_result_page_shows_stages(self):
        response = self.client.get(reverse("run_detail_v2", args=[self.run.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Этапы конфигурации тормозов")
        self.assertContains(response, "итерационный · этапов: 2")
        self.assertContains(response, "rs-changed")
        self.assertContains(response, reverse("iterative_detail", args=[self.calc.pk]))

        wb = openpyxl.load_workbook(self.run.report_file.path, read_only=True)
        self.assertIn("Этапы", wb.sheetnames)
        wb.close()

    def test_results_list_badge(self):
        response = self.client.get(reverse("results"))
        self.assertContains(response, "итерац.")
